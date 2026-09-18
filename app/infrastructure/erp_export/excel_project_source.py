from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ...application.errors import ApplicationOperationError
from ...application.project_sync import ExternalProjectRecord, ProjectSourcePort


_REQUIRED_COLUMNS = (
    "ID projet",
    "Statut",
    "Description",
    "Nom du client",
    "Gestionnaire de projet",
)

_STYLES_PATH = "xl/styles.xml"
_SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_OPENPYXL_EMPTY_FILL_ERROR = "openpyxl.styles.fills.Fill"


def _text(value: object) -> str:
    return str(value or "").strip()


def _optional_text(value: object) -> str | None:
    text = _text(value)
    return text or None


def _project_number(value: object) -> str:
    if isinstance(value, bool):
        return _text(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _workbook_with_repaired_empty_fills(path: Path) -> BytesIO | None:
    """Return an in-memory XLSX with invalid empty fill entries repaired.

    Some ERP Excel exports contain empty fill records in xl/styles.xml. Excel
    tolerates them, but openpyxl 3.1.x rejects them with an expected Fill
    TypeError. The importer only reads values, so replacing an empty fill with
    a neutral patternFill is safe and keeps the same fill index.
    """

    try:
        with ZipFile(path, "r") as source:
            if _STYLES_PATH not in source.namelist():
                return None

            root = ElementTree.fromstring(source.read(_STYLES_PATH))
            fills = root.find(f"{{{_SPREADSHEET_NS}}}fills")
            if fills is None:
                return None

            repaired = 0
            for fill in fills.findall(f"{{{_SPREADSHEET_NS}}}fill"):
                if len(fill) == 0:
                    ElementTree.SubElement(
                        fill,
                        f"{{{_SPREADSHEET_NS}}}patternFill",
                        {"patternType": "none"},
                    )
                    repaired += 1

            if repaired == 0:
                return None

            ElementTree.register_namespace("", _SPREADSHEET_NS)
            repaired_styles = ElementTree.tostring(
                root,
                encoding="utf-8",
                xml_declaration=True,
            )

            buffer = BytesIO()
            with ZipFile(buffer, "w") as target:
                for item in source.infolist():
                    content = (
                        repaired_styles
                        if item.filename == _STYLES_PATH
                        else source.read(item.filename)
                    )
                    target.writestr(item, content)
    except (BadZipFile, ElementTree.ParseError, OSError):
        return None

    buffer.seek(0)
    return buffer


def _load_erp_workbook(path: Path):
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except TypeError as exc:
        if _OPENPYXL_EMPTY_FILL_ERROR not in str(exc):
            raise

        repaired = _workbook_with_repaired_empty_fills(path)
        if repaired is None:
            raise

        return load_workbook(repaired, read_only=True, data_only=True)


class ErpExcelProjectSource(ProjectSourcePort):
    """Project source for the manual ERP XLSX export used before live ERP connectivity."""

    def __init__(self, path: str | Path, *, sheet_name: str = "Données") -> None:
        self._path = Path(path)
        self._sheet_name = sheet_name

    @property
    def path(self) -> Path:
        return self._path

    def _header_indexes(self, header: Iterable[object]) -> dict[str, int]:
        indexes: dict[str, int] = {}
        for index, value in enumerate(header):
            label = _text(value)
            if label and label not in indexes:
                indexes[label] = index

        missing = [column for column in _REQUIRED_COLUMNS if column not in indexes]
        if missing:
            raise ApplicationOperationError(
                "Le fichier d'export ERP ne contient pas toutes les colonnes requises.",
                code="erp_excel_project_headers_invalid",
                context={"missing_columns": missing, "sheet": self._sheet_name},
            )
        return indexes

    @staticmethod
    def _value(row: tuple[object, ...], index: int) -> object | None:
        return row[index] if index < len(row) else None

    def list_projects(self) -> tuple[ExternalProjectRecord, ...]:
        if not self._path.is_file():
            raise ApplicationOperationError(
                "Le fichier d'export ERP est introuvable.",
                code="erp_excel_project_file_not_found",
                context={"path": str(self._path)},
            )

        try:
            workbook = _load_erp_workbook(self._path)
        except (OSError, InvalidFileException, TypeError, ValueError) as exc:
            raise ApplicationOperationError(
                "Impossible d'ouvrir le fichier d'export ERP.",
                code="erp_excel_project_file_invalid",
                context={"path": str(self._path)},
            ) from exc

        try:
            if self._sheet_name not in workbook.sheetnames:
                raise ApplicationOperationError(
                    f"La feuille '{self._sheet_name}' est absente du fichier d'export ERP.",
                    code="erp_excel_project_sheet_missing",
                    context={"sheet": self._sheet_name},
                )

            worksheet = workbook[self._sheet_name]
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                raise ApplicationOperationError(
                    "La feuille des projets ERP est vide.",
                    code="erp_excel_project_sheet_empty",
                    context={"sheet": self._sheet_name},
                )
            indexes = self._header_indexes(header)

            result: list[ExternalProjectRecord] = []
            seen_numbers: set[str] = set()
            for row_number, row in enumerate(rows, start=2):
                if not any(value not in (None, "") for value in row):
                    continue

                number = _project_number(self._value(row, indexes["ID projet"]))
                name = _text(self._value(row, indexes["Description"]))
                if not number or not name:
                    raise ApplicationOperationError(
                        "Une ligne de l'export ERP contient un projet incomplet.",
                        code="erp_excel_project_row_invalid",
                        context={
                            "row": row_number,
                            "has_project_id": bool(number),
                            "has_description": bool(name),
                        },
                    )
                if number in seen_numbers:
                    raise ApplicationOperationError(
                        "Le même numéro de projet apparaît plus d'une fois dans l'export ERP.",
                        code="erp_excel_project_duplicate_number",
                        context={"row": row_number, "project_number": number},
                    )
                seen_numbers.add(number)

                result.append(
                    ExternalProjectRecord(
                        # The XLSX export exposes ProjectID, not Acumatica's REST row id.
                        # Keep the external id unbound so a later live ERP sync can adopt
                        # the same local project by number and attach the real ERP id.
                        external_id=None,
                        number=number,
                        name=name,
                        client=_optional_text(self._value(row, indexes["Nom du client"])),
                        project_manager_name=_optional_text(
                            self._value(row, indexes["Gestionnaire de projet"])
                        ),
                        status=_text(self._value(row, indexes["Statut"])) or "Actif",
                    )
                )
            return tuple(result)
        finally:
            workbook.close()
