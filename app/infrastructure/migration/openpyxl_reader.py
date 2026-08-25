from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


class CutoverSourceError(RuntimeError):
    pass


class OpenpyxlCutoverReader:
    """Read a frozen xlsx/xlsm directly from disk without opening Excel or saving it."""

    def __init__(self, workbook_path: str | Path) -> None:
        self.path = Path(workbook_path).expanduser().resolve()
        if not self.path.exists():
            raise CutoverSourceError(f"Classeur introuvable: {self.path}")
        if self.path.suffix.casefold() not in {".xlsx", ".xlsm"}:
            raise CutoverSourceError("Le cutover exige un classeur .xlsx ou .xlsm.")
        try:
            self._workbook = load_workbook(
                filename=self.path,
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        except Exception as exc:
            raise CutoverSourceError(
                f"Impossible d'ouvrir le classeur en lecture seule: {self.path}"
            ) from exc

    def close(self) -> None:
        self._workbook.close()

    def __enter__(self) -> "OpenpyxlCutoverReader":
        return self

    def __exit__(self, exc_type: object, exc: object, traceback: object) -> None:
        self.close()

    @staticmethod
    def _header_index(values: tuple[Any, ...], expected_header: str) -> int | None:
        expected = str(expected_header).strip()
        for index, value in enumerate(values):
            if str(value or "").strip() == expected:
                return index
        return None

    def records(self, sheet: str, expected_header: str) -> tuple[dict[str, Any], ...]:
        if sheet not in self._workbook.sheetnames:
            raise CutoverSourceError(f"Feuille requise absente: {sheet}")
        worksheet = self._workbook[sheet]

        header_row: int | None = None
        headers: tuple[Any, ...] | None = None
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=1, max_row=20, values_only=True),
            start=1,
        ):
            row_values = tuple(values)
            if self._header_index(row_values, expected_header) is not None:
                header_row = row_number
                headers = row_values
                break
        if header_row is None or headers is None:
            raise CutoverSourceError(
                f"En-tête '{expected_header}' introuvable dans les 20 premières lignes de {sheet}."
            )

        keys: list[str | None] = []
        seen: dict[str, int] = {}
        for index, raw in enumerate(headers, start=1):
            text = str(raw or "").strip()
            if not text:
                keys.append(None)
                continue
            seen[text] = seen.get(text, 0) + 1
            key = text if seen[text] == 1 else f"{text} [{get_column_letter(index)}]"
            keys.append(key)

        result: list[dict[str, Any]] = []
        for excel_row, values in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            row_values = tuple(values)
            if not any(value not in (None, "") for value in row_values):
                continue
            item: dict[str, Any] = {"_row": excel_row}
            for index, key in enumerate(keys):
                if key is None:
                    continue
                item[key] = row_values[index] if index < len(row_values) else None
            result.append(item)
        return tuple(result)
